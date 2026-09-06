"""Extract plain text from uploaded CV files (PDF, DOCX, plain text)."""

from __future__ import annotations

import io
from pathlib import Path

from app.core.errors import ApiError

MAX_CV_BYTES = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}
ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",  # rejected later if .doc (legacy); kept for clients that mislabel
    "text/plain",
    "text/markdown",
    "text/csv",
    "application/octet-stream",
}


def _normalize_text(text: str) -> str:
    cleaned = text.replace("\r\n", "\n").replace("\r", "\n")
    # Collapse excessive blank lines from PDF extractors.
    while "\n\n\n" in cleaned:
        cleaned = cleaned.replace("\n\n\n", "\n\n")
    return cleaned.strip()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    return _normalize_text("\n".join(pages))


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _normalize_text("\n".join(parts))


def _extract_plain(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return _normalize_text(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    return _normalize_text(data.decode("utf-8", errors="ignore"))


def extract_cv_text(
    *,
    data: bytes,
    filename: str | None = None,
    content_type: str | None = None,
) -> tuple[str, str]:
    """Return (text, detected_format) for a CV upload.

    Raises ApiError for unsupported/empty/oversized files.
    """
    if len(data) > MAX_CV_BYTES:
        raise ApiError(
            status_code=413,
            code="cv_file_too_large",
            message="CV files must be 5 MB or smaller.",
        )
    if not data:
        raise ApiError(
            status_code=422,
            code="cv_file_empty",
            message="The uploaded CV file is empty.",
        )

    name = (filename or "cv.txt").strip() or "cv.txt"
    suffix = Path(name).suffix.lower()
    ctype = (content_type or "").split(";")[0].strip().lower()

    if suffix == ".doc":
        raise ApiError(
            status_code=422,
            code="cv_format_unsupported",
            message="Legacy .doc files are not supported. Please upload PDF, DOCX, or plain text.",
        )

    if suffix and suffix not in ALLOWED_EXTENSIONS:
        raise ApiError(
            status_code=422,
            code="cv_format_unsupported",
            message="Supported CV formats: PDF, DOCX, or plain text (.txt).",
        )

    if ctype and ctype not in ALLOWED_CONTENT_TYPES and suffix not in ALLOWED_EXTENSIONS:
        raise ApiError(
            status_code=422,
            code="cv_format_unsupported",
            message="Supported CV formats: PDF, DOCX, or plain text (.txt).",
        )

    try:
        if suffix == ".pdf" or ctype == "application/pdf":
            text = _extract_pdf(data)
            fmt = "pdf"
        elif suffix == ".docx" or ctype.endswith(
            "wordprocessingml.document"
        ):
            text = _extract_docx(data)
            fmt = "docx"
        else:
            text = _extract_plain(data)
            fmt = "text"
    except ApiError:
        raise
    except Exception as exc:
        raise ApiError(
            status_code=422,
            code="cv_extract_failed",
            message=f"Could not read text from this {suffix or 'file'}. Try PDF, DOCX, or plain text.",
        ) from exc

    if len(text) < 40:
        raise ApiError(
            status_code=422,
            code="cv_text_too_short",
            message=(
                "Not enough readable text was found in this file. "
                "Scanned image PDFs are not supported yet — use a text PDF, DOCX, or paste the text."
            ),
        )
    if len(text) > 100_000:
        text = text[:100_000]

    return text, fmt
